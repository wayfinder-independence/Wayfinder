"""Schedule Reader v0: read a weekly schedule from a Word (.docx) file into JSON.

The reader expects the layout people commonly type by hand in Word:

* some lines of text above a table (a title, a theme, ...);
* one table whose first row holds day headings such as "Mon Mar 1";
* following rows that each start with a period label such as "MORNING";
* in each day/period cell, numbered activities ("1. Yoga") each followed by a
  line listing who is assigned, plus any other lines (notes, "Closed", ...).

Everything the document says is kept as source text. Anything Wayfinder works
out for itself (dates, activity profiles) is labelled as such, and anything it
cannot work out is listed under "needs_review" instead of being guessed.

Usage:
    python -m wayfinder.schedule_reader SCHEDULE.docx [-o OUTPUT.json]
"""

import argparse
import calendar
import datetime
import json
import re
import sys
from pathlib import Path

import docx
from docx.document import Document

from wayfinder.activity_profile import suggest_profile

FORMAT_VERSION = "0.1"
REPOSITORY_ROOT = Path(__file__).resolve().parent.parent

WEEKDAYS = {"mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6}
MONTHS = {name.lower(): number for number, name in enumerate(calendar.month_abbr) if name}
PERIODS = ("morning", "afternoon", "evening", "lunch", "all day")

DAY_HEADING = re.compile(
    r"^(?P<weekday>mon|tue|wed|thu|fri|sat|sun)[a-z]*\.?,?\s+"
    r"(?P<month>[a-z]{3})[a-z]*\.?\s+(?P<day>\d{1,2})(?:st|nd|rd|th)?(?:,?\s+(?P<year>\d{4}))?$",
    re.IGNORECASE,
)
NUMBERED_ACTIVITY = re.compile(r"^(?P<number>\d{1,2})\s*[.)]\s*(?P<name>.+)$")
CLOSURE_WORDS = re.compile(r"\b(closed|cancell?ed|no program)\b", re.IGNORECASE)
YEAR = re.compile(r"\b(20\d\d)\b")


class ScheduleReadError(Exception):
    """The file could not be read as a weekly schedule."""


# ---------------------------------------------------------------- text helpers

def _clean(text):
    """Collapse runs of spaces/tabs; keep the words exactly as written."""
    return re.sub(r"\s+", " ", text).strip()


def _is_bold(paragraph):
    """True when every visible piece of text in the paragraph is bold."""
    visible = [run for run in paragraph.runs if run.text.strip()]
    if not visible:
        return False
    style_bold = bool(paragraph.style is not None and paragraph.style.font.bold)
    return all(run.bold or (run.bold is None and style_bold) for run in visible)


def _cell_lines(cell):
    """Non-empty lines of a table cell as (text, is_bold) pairs."""
    lines = []
    for paragraph in cell.paragraphs:
        for piece in paragraph.text.split("\n"):
            text = _clean(piece)
            if text:
                lines.append((text, _is_bold(paragraph)))
    return lines


# ---------------------------------------------------------------- days and dates

def _parse_day_heading(text, fallback_year):
    """Turn "Mon Mar 1" into weekday/date facts. Returns (day dict, problem or None)."""
    day = {"heading": text, "weekday": None, "date": None}
    match = DAY_HEADING.match(text)
    if not match:
        return day, f'The day heading "{text}" could not be understood as a weekday and date.'
    weekday_number = WEEKDAYS[match["weekday"][:3].lower()]
    day["weekday"] = calendar.day_name[weekday_number]
    month = MONTHS.get(match["month"][:3].lower())
    year = int(match["year"]) if match["year"] else fallback_year
    if month is None:
        return day, f'The month in the day heading "{text}" was not recognised.'
    if year is None:
        return day, f'The day heading "{text}" has no year and no year was found elsewhere, so no full date is given.'
    try:
        date = datetime.date(year, month, int(match["day"]))
    except ValueError:
        return day, f'The day heading "{text}" is not a real calendar date.'
    if date.weekday() != weekday_number:
        return day, (f'The day heading "{text}" says {day["weekday"]}, but in {year} that date is a '
                     f'{calendar.day_name[date.weekday()]}. The date was left blank.')
    day["date"] = date.isoformat()
    return day, None


def _find_year(title_lines, file_name):
    for text, source in [(line, "the document text") for line in title_lines] + [(file_name, "the file name")]:
        found = YEAR.search(text or "")
        if found:
            return int(found.group(1)), source
    return None, None


def _find_schedule_table(document):
    for table in document.tables:
        if not table.rows:
            continue
        headings = [_clean(cell.text) for cell in table.rows[0].cells[1:]]
        if any(DAY_HEADING.match(h) for h in headings):
            return table
    return None


# ---------------------------------------------------------------- one cell

def _split_participants(text):
    return [_clean(name) for name in text.split(",") if _clean(name)]


def _read_cell(lines, where):
    """Read one day/period cell into activities, notes and a status."""
    activities, notes, review = [], [], []
    current = None
    for text, bold in lines:
        numbered = NUMBERED_ACTIVITY.match(text)
        if numbered:
            current = {
                "id": f"{where}-{len(activities) + 1}",
                "number": int(numbered["number"]),
                "name": _clean(numbered["name"]),
                "source_text": text,
                "participants": None,
            }
            activities.append(current)
        elif current is not None and current["participants"] is None and not bold:
            # The first plain (non-bold) line after an activity lists who is assigned.
            current["participants"] = {"source_text": text, "names": _split_participants(text)}
        else:
            notes.append({"text": text, "follows_activity": current["number"] if current else None})

    for activity in activities:
        if activity["participants"] is None:
            review.append({"where": activity["id"], "issue": "No list of participants was found for this activity."})
        activity["classification"] = suggest_profile(activity["name"])

    all_text = " ".join(text for text, _ in lines)
    if activities:
        status = "activities"
        if CLOSURE_WORDS.search(all_text):
            review.append({"where": where, "issue": "This period lists activities but also mentions a closure or cancellation."})
    elif CLOSURE_WORDS.search(all_text):
        status = "closed"
    elif not lines:
        status = "empty"
        review.append({"where": where, "issue": "Nothing is written for this period."})
    else:
        status = "unclear"
        review.append({"where": where, "issue": (
            f'No activities are listed. The schedule only says: "{all_text}". '
            "It is not clear whether the program runs at this time.")})

    for note in notes:
        if note["follows_activity"] is not None and activities:
            review.append({"where": where, "issue": (
                f'The line "{note["text"]}" appears after activity {note["follows_activity"]}. '
                "It may belong to that activity or to the whole period.")})
    return {"status": status, "activities": activities, "notes": notes}, review


# ---------------------------------------------------------------- whole document

def read_schedule(path):
    """Read a .docx weekly schedule and return the Wayfinder schedule as a dict."""
    path = Path(path)
    try:
        document = docx.Document(str(path))
    except Exception as error:
        raise ScheduleReadError(
            f"{path.name} could not be opened as a Word .docx file. If it is protected with "
            "restricted access (rights management) or is an old .doc file, open it in Word and "
            "save an unprotected .docx copy outside the Wayfinder folder."
        ) from error
    return read_document(document, path.name)


def read_document(document: Document, file_name=None):
    review = []
    table = _find_schedule_table(document)
    if table is None:
        raise ScheduleReadError("No table with day headings (for example \"Mon Mar 1\") was found.")

    # Text outside the table, kept in order.
    before, after, seen_table = [], [], False
    for block in document.iter_inner_content():
        if hasattr(block, "rows"):
            if block._tbl is table._tbl:
                seen_table = True
            else:
                review.append({"where": "document", "issue": "The document has another table that v0 does not read."})
            continue
        text = _clean(block.text)
        if text:
            (after if seen_table else before).append(text)

    year, year_source = _find_year(before, file_name)
    days, columns = [], []
    for column, cell in enumerate(table.rows[0].cells[1:], start=1):
        day, problem = _parse_day_heading(_clean(cell.text), year)
        day["periods"] = []
        days.append(day)
        columns.append(column)
        if problem:
            review.append({"where": f"day{column}", "issue": problem})
    missing_year = [r for r in review if "has no year" in r["issue"]]
    if missing_year:
        review = [r for r in review if r not in missing_year]
        review.append({"where": ", ".join(r["where"] for r in missing_year), "issue": (
            "The day headings have no year and no year was found in the title or file name, "
            "so no full dates are given.")})

    for row in table.rows[1:]:
        cells = row.cells
        label_lines = [text for text, _ in _cell_lines(cells[0])]
        label = " ".join(label_lines)
        first = label_lines[0].lower() if label_lines else ""
        period = next((p for p in PERIODS if first.startswith(p)), None)
        period_key = (period or f"row{len(days[0]['periods']) + 1}").replace(" ", "-")
        if period is None:
            review.append({"where": period_key, "issue": f'The row label "{label}" is not a recognised period such as morning or afternoon.'})
        for day, column in zip(days, columns):
            where = f"day{column}-{period_key}"
            if column < len(cells) and column > 1 and cells[column]._tc is cells[column - 1]._tc:
                review.append({"where": where, "issue": "This cell is merged across more than one day; its contents were repeated for each day."})
            lines = _cell_lines(cells[column]) if column < len(cells) else []
            content, cell_review = _read_cell(lines, where)
            review.extend(cell_review)
            day["periods"].append({"id": where, "label": label or None, "period": period, **content})

    review.extend(_participant_notation_review(days))
    for day in days:
        for slot in day["periods"]:
            for activity in slot["activities"]:
                # Anything short of high confidence (medium, low or unknown) is shown to a person.
                to_check = []
                for dimension in ("environment", "intensity"):
                    suggestion = activity["classification"][dimension]
                    if suggestion["suggested"] == "unknown":
                        to_check.append(f"{dimension} (unknown)")
                    elif suggestion["confidence"] != "high":
                        to_check.append(f'{dimension} ({suggestion["confidence"]} confidence)')
                if to_check:
                    review.append({"where": activity["id"], "issue": (
                        f'Please check Wayfinder\'s suggested {" and ".join(to_check)} for "{activity["name"]}".')})

    dates = [day["date"] for day in days if day["date"]]
    return {
        "format": "wayfinder-weekly-schedule",
        "format_version": FORMAT_VERSION,
        "source": {"file_name": file_name},
        "title": before[0] if before else None,
        "other_text_before_table": before[1:],
        "text_after_table": after,
        "date_range": {
            "start": min(dates) if dates else None,
            "end": max(dates) if dates else None,
            "worked_out_by_wayfinder": True,
            "basis": ("From the day headings in the table"
                      + (f"; the year comes from {year_source}." if year_source else "; no year was found.")),
        },
        "days": days,
        "needs_review": review,
        "corrections": [],
    }


def _participant_notation_review(days):
    """Flag name entries whose notation the schedule does not explain, once per distinct entry."""
    found = {}
    for day in days:
        for slot in day["periods"]:
            for activity in slot["activities"]:
                for name in (activity["participants"] or {}).get("names", []):
                    if "(" in name or re.search(r"[\w.]-\w", name):
                        found.setdefault(name, []).append(activity["id"])
    review = []
    for name, places in found.items():
        meaning = ("is written in brackets" if "(" in name
                   else "joins names with a hyphen (a double-barrelled name or a pairing?)")
        review.append({"where": ", ".join(places),
                       "issue": f'The entry "{name}" {meaning}; the schedule does not say what this means.'})
    return review


# ---------------------------------------------------------------- command line

def main(argv=None):
    parser = argparse.ArgumentParser(description="Read a weekly schedule (.docx) into Wayfinder JSON.")
    parser.add_argument("schedule", help="path to the .docx schedule")
    parser.add_argument("-o", "--output", help="where to write the JSON (default: print it)")
    args = parser.parse_args(argv)

    try:
        result = read_schedule(args.schedule)
    except ScheduleReadError as error:
        print(f"Could not read the schedule: {error}", file=sys.stderr)
        return 1

    text = json.dumps(result, indent=2, ensure_ascii=False)
    if not args.output:
        print(text)
        return 0
    output = Path(args.output).resolve()
    if REPOSITORY_ROOT in output.parents:
        print("Refusing to write schedule output inside the Wayfinder folder, because real schedules "
              "are private. Choose a location outside it.", file=sys.stderr)
        return 1
    output.write_text(text + "\n", encoding="utf-8")
    print(f"Wrote {output} ({len(result['needs_review'])} items need review).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
