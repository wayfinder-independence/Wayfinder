"""Tests for Schedule Reader v0.

All schedule content here is invented for testing. The documents are built in
code at test time, so no Word files are stored in the repository.
"""

import json

import docx
import pytest

from wayfinder.schedule_reader import ScheduleReadError, main, read_document, read_schedule

# An invented week at an invented day program.
DAYS = ["Mon Mar 1", "Tue Mar 2", "Wed  Mar 3"]
CELLS = {
    "MORNING": [
        [("1. Spring Hike", True), ("Alex, Priya, Jordan", False),
         ("2. Baking Club", True), ("Sam, Lee", False)],
        [("CENTRE", True), ("Closed for staff training", True)],
        [("1. Pottery Wheel", True), ("Alex, Sam", False),
         ("2. Mystery Outing", True)],
    ],
    "AFTERNOON": [
        [("1. Library Visit", True), ("Alex, (Robin), Jordan-Casey, Kim T.-Ray", False),
         ("Bring your library card", True)],
        [("Enjoy your day off!", True)],
        [],
    ],
}


def build_document(day_headings=DAYS, title="RIVERSIDE PROGRAM WEEK OF MARCH 1 2027"):
    document = docx.Document()
    document.add_paragraph().add_run(title).bold = True
    document.add_paragraph("Theme: Spring is here")
    table = document.add_table(rows=1 + len(CELLS), cols=1 + len(day_headings))
    for column, heading in enumerate(day_headings, start=1):
        table.rows[0].cells[column].text = heading
    for row_number, (label, cells) in enumerate(CELLS.items(), start=1):
        table.rows[row_number].cells[0].text = f"{label}\nGroup time"
        for column, lines in enumerate(cells, start=1):
            cell = table.rows[row_number].cells[column]
            cell.paragraphs[0].text = ""
            for text, bold in lines:
                cell.add_paragraph().add_run(text).bold = bold
    document.add_paragraph("Questions? Call the front desk.")
    return document


@pytest.fixture
def schedule():
    return read_document(build_document(), "invented.docx")


def period(schedule, day_index, name):
    return next(p for p in schedule["days"][day_index]["periods"] if p["period"] == name)


def test_title_and_text_outside_the_table_are_kept(schedule):
    assert schedule["title"] == "RIVERSIDE PROGRAM WEEK OF MARCH 1 2027"
    assert schedule["other_text_before_table"] == ["Theme: Spring is here"]
    assert schedule["text_after_table"] == ["Questions? Call the front desk."]


def test_days_and_date_range(schedule):
    assert [d["weekday"] for d in schedule["days"]] == ["Monday", "Tuesday", "Wednesday"]
    assert [d["date"] for d in schedule["days"]] == ["2027-03-01", "2027-03-02", "2027-03-03"]
    assert schedule["days"][2]["heading"] == "Wed Mar 3"
    assert schedule["date_range"]["start"] == "2027-03-01"
    assert schedule["date_range"]["end"] == "2027-03-03"
    assert schedule["date_range"]["worked_out_by_wayfinder"] is True


def test_morning_and_afternoon_with_several_activities(schedule):
    monday = schedule["days"][0]
    assert [p["period"] for p in monday["periods"]] == ["morning", "afternoon"]
    assert monday["periods"][0]["label"] == "MORNING Group time"
    morning = period(schedule, 0, "morning")
    assert [a["name"] for a in morning["activities"]] == ["Spring Hike", "Baking Club"]
    assert morning["activities"][0]["source_text"] == "1. Spring Hike"
    assert morning["activities"][1]["participants"] == {"source_text": "Sam, Lee", "names": ["Sam", "Lee"]}


def test_closed_period_is_recorded_as_closed(schedule):
    closed = period(schedule, 1, "morning")
    assert closed["status"] == "closed"
    assert closed["activities"] == []
    assert [n["text"] for n in closed["notes"]] == ["CENTRE", "Closed for staff training"]


def test_text_without_activities_is_kept_and_flagged_not_guessed(schedule):
    afternoon = period(schedule, 1, "afternoon")
    assert afternoon["status"] == "unclear"
    assert afternoon["notes"][0]["text"] == "Enjoy your day off!"
    assert any(r["where"] == afternoon["id"] for r in schedule["needs_review"])


def test_empty_period_is_flagged(schedule):
    assert period(schedule, 2, "afternoon")["status"] == "empty"


def test_missing_participants_are_null_and_flagged(schedule):
    outing = period(schedule, 2, "morning")["activities"][1]
    assert outing["name"] == "Mystery Outing"
    assert outing["participants"] is None
    assert any(r["where"] == outing["id"] and "participants" in r["issue"] for r in schedule["needs_review"])


def test_note_after_activity_keeps_its_position_and_is_flagged(schedule):
    afternoon = period(schedule, 0, "afternoon")
    assert afternoon["notes"] == [{"text": "Bring your library card", "follows_activity": 1}]
    assert any("Bring your library card" in r["issue"] for r in schedule["needs_review"])


def test_unexplained_name_notation_is_kept_and_flagged(schedule):
    names = period(schedule, 0, "afternoon")["activities"][0]["participants"]["names"]
    assert names == ["Alex", "(Robin)", "Jordan-Casey", "Kim T.-Ray"]
    issues = " ".join(r["issue"] for r in schedule["needs_review"])
    assert '"(Robin)"' in issues and '"Jordan-Casey"' in issues and '"Kim T.-Ray"' in issues


def test_suggestions_are_separate_from_confirmation(schedule):
    hike = period(schedule, 0, "morning")["activities"][0]["classification"]
    assert hike["environment"]["suggested"] == "outdoor"
    assert hike["environment"]["confirmed"] is None
    assert hike["intensity"]["confirmed"] is None
    assert hike["review_status"] == "pending"
    assert schedule["corrections"] == []


def test_ambiguous_activity_is_unknown_and_flagged(schedule):
    pottery = period(schedule, 2, "morning")["activities"][0]
    assert pottery["classification"]["environment"]["suggested"] == "unknown"
    assert pottery["classification"]["intensity"]["suggested"] == "unknown"
    assert any(r["where"] == pottery["id"] and "environment (unknown) and intensity (unknown)" in r["issue"]
               for r in schedule["needs_review"])


def test_medium_confidence_suggestion_is_flagged_but_high_confidence_is_not(schedule):
    # Library Visit: environment is high confidence, intensity is medium confidence.
    library = period(schedule, 0, "afternoon")["activities"][0]
    assert library["classification"]["environment"]["confidence"] == "high"
    assert library["classification"]["intensity"]["confidence"] == "medium"
    issues = [r["issue"] for r in schedule["needs_review"]
              if r["where"] == library["id"] and "suggested" in r["issue"]]
    assert issues == ['Please check Wayfinder\'s suggested intensity (medium confidence) for "Library Visit".']


def test_activity_with_only_high_confidence_suggestions_is_not_flagged():
    document = docx.Document()
    document.add_paragraph("GARDEN WEEK 2027")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[1].text = "Mon Mar 1"
    table.rows[1].cells[0].text = "MORNING"
    cell = table.rows[1].cells[1]
    cell.paragraphs[0].add_run("1. Garden Yoga").bold = True
    cell.add_paragraph("Alex")
    schedule = read_document(document, "garden.docx")
    profile = schedule["days"][0]["periods"][0]["activities"][0]["classification"]
    assert profile["environment"]["confidence"] == profile["intensity"]["confidence"] == "high"
    assert schedule["needs_review"] == []


def test_no_year_means_no_invented_date():
    schedule = read_document(build_document(title="RIVERSIDE PROGRAM"), "no-year.docx")
    assert schedule["days"][0]["weekday"] == "Monday"
    assert schedule["days"][0]["date"] is None
    assert schedule["date_range"]["start"] is None
    assert len([r for r in schedule["needs_review"] if "no year" in r["issue"]]) == 1


def test_weekday_that_does_not_match_the_date_is_not_trusted():
    schedule = read_document(build_document(day_headings=["Fri Mar 1", "Tue Mar 2", "Wed Mar 3"]), "x.docx")
    assert schedule["days"][0]["date"] is None
    assert any("Friday" in r["issue"] for r in schedule["needs_review"])


def test_file_that_is_not_a_docx_gives_a_clear_error(tmp_path):
    fake = tmp_path / "schedule.docx"
    fake.write_bytes(b"not a word document")
    with pytest.raises(ScheduleReadError, match="could not be opened"):
        read_schedule(fake)


def test_command_line_writes_json_outside_the_repository(tmp_path):
    source, output = tmp_path / "week.docx", tmp_path / "week.json"
    build_document().save(source)
    assert main([str(source), "-o", str(output)]) == 0
    assert json.loads(output.read_text(encoding="utf-8"))["days"][0]["weekday"] == "Monday"


def test_command_line_refuses_to_write_inside_the_repository(tmp_path):
    source = tmp_path / "week.docx"
    build_document().save(source)
    inside = __import__("pathlib").Path(__file__).resolve().parent / "should-not-exist.json"
    assert main([str(source), "-o", str(inside)]) == 1
    assert not inside.exists()
